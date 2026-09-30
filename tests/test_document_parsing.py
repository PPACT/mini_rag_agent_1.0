"""解析层（2.0-2 路由 / 2.0-1 表格 / 2.0-3 显式失败）的测试。

夹具**现场生成**（python-docx / openpyxl 造），所以不依赖任何外部语料 ——
真实文档到位后，这里再加一层"对真实语料"的验收（见 `2.0-6`）。
"""
from __future__ import annotations

import pytest

from src.document_parser.loader import load_document, load_text
from src.document_parser.routing import (
    ParseFailedError,
    UnsupportedFormatError,
    parse_document,
    route_file,
    supported_extensions,
)
from src.document_parser.tables import missing_numbers, render_nl


# ---------------------------------------------------------------- 夹具


def _make_docx(path, *, paragraphs: list[str], table: list[list[str]] | None = None):
    import docx

    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, val in enumerate(row):
                t.cell(i, j).text = str(val)
    doc.save(str(path))
    return path


def _make_xlsx(path, *, sheets: dict[str, list[list[str]]]):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for r in rows:
            ws.append(r)
    wb.save(str(path))
    return path


REIMBURSE_TABLE = [
    ["城市等级", "住宿上限（元/晚）"],
    ["一线城市", "600"],
    ["二线城市", "400"],
    ["其他城市", "300"],
]


# ---------------------------------------------------------------- 2.0-3 显式失败


def test_known_unsupported_says_why(tmp_path):
    """`.doc` 认得出但本轮不做 → **明说原因**，不是含糊的"不支持"。"""
    p = tmp_path / "a.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 20)   # OLE2 头
    with pytest.raises(UnsupportedFormatError) as ei:
        route_file(str(p))
    assert "旧版二进制" in str(ei.value)


def test_all_parsers_failed_reports_each_attempt(tmp_path):
    """格式认得、但解析器全失败 → 报错里要有**每个解析器各自的原因**（2.0-3）。"""
    p = tmp_path / "broken.pdf"
    # ⚠️ bytes 字面量只能是 ASCII —— 中文要显式 encode（本节就是被这条绊了一跤）
    p.write_bytes(b"%PDF-1.4\n" + "这不是真的 PDF 结构".encode())
    with pytest.raises(ParseFailedError) as ei:
        parse_document(str(p))
    detail = ei.value.detail
    assert "pdfplumber" in str(detail["attempts"])
    assert detail["route"]           # 路由决策也要留下来，便于定位
    assert ei.value.path == str(p)


def test_failure_is_not_a_silent_empty(tmp_path):
    """⚠️ 关键回归：**不能"解析不出内容"却当成功返回空**。

    旧实现只抛一句 ValueError；现在的失败必须带文件、路由、逐解析器原因。
    """
    p = tmp_path / "broken.docx"
    p.write_bytes(b"not a zip at all")
    with pytest.raises(ParseFailedError):
        parse_document(str(p))


# ---------------------------------------------------------------- 2.0-2 文件级路由


@pytest.mark.parametrize("name,kind", [
    ("a.pdf", "pdf"), ("a.docx", "docx"), ("a.xlsx", "xlsx"), ("a.pptx", "pptx"),
    ("a.csv", "csv"), ("a.md", "markdown"), ("a.txt", "text"), ("a.py", "code"),
])
def test_route_by_extension(tmp_path, name, kind):
    p = tmp_path / name
    p.write_bytes(b"x")
    plan = route_file(str(p))
    assert plan.kind == kind
    assert plan.confidence == "high"
    assert any(s.startswith("后缀.") for s in plan.signals)


def test_route_by_magic_when_extension_lies(tmp_path):
    """后缀说是 txt，内容却是 PDF → **magic number 兜住**，且信号可查（P-2）。"""
    p = tmp_path / "fake.txt"
    p.write_bytes(b"%PDF-1.4\n%...")
    plan = route_file(str(p))
    assert plan.kind == "pdf"
    assert any(s.startswith("magic.") for s in plan.signals)


def test_unknown_extension_falls_back_to_text(tmp_path):
    p = tmp_path / "a.weird"
    p.write_bytes("你好".encode())
    plan = route_file(str(p))
    assert plan.kind == "text"
    assert plan.confidence == "low"
    assert "兜底.纯文本" in plan.signals


def test_supported_extensions_is_single_source(tmp_path):
    """上传接口引用的那份清单必须**来自路由表**（不再各写一份）。"""
    exts = supported_extensions()
    assert {"pdf", "docx", "xlsx", "pptx", "csv", "md", "txt", "py"} <= set(exts)
    assert "doc" not in exts and "ppt" not in exts, "旧版二进制格式从来读不了，不该放行"


# ---------------------------------------------------------------- 2.0-1 表格入库


def test_docx_table_becomes_a_table_block_and_keeps_numbers(tmp_path):
    """**2.0-1 的验收（文档级）**：表格里的每个数值都要能被找到 —— 且表格不再被丢掉。"""
    p = _make_docx(tmp_path / "d.docx",
                   paragraphs=["# 差旅制度", "本制度适用于全体员工。"],
                   table=REIMBURSE_TABLE)

    doc = load_document(str(p))
    tables = doc.tables()
    assert len(tables) == 1, "docx 里的表格必须被提出来（旧实现整块丢掉）"
    tb = tables[0]
    assert tb.raw_table and "| 一线城市 | 600 |" in tb.raw_table
    assert "住宿上限（元/晚）=600" in tb.text
    # 验收闸：原表里每个数值都在自然语言版里
    assert missing_numbers(tb.raw_table, tb.text) == []
    # 整份文档的纯文本里也能找到（表格不再消失在 load_text 的输出里）
    assert "600" in load_text(str(p))


def test_docx_preserves_document_order(tmp_path):
    """段落与表格必须按**文档顺序**产出 —— 表格在第几段之后决定切块边界。"""
    p = _make_docx(tmp_path / "o.docx", paragraphs=["前一段"], table=REIMBURSE_TABLE)
    import docx

    doc = docx.Document(str(p))
    doc.add_paragraph("后一段")
    doc.save(str(p))

    kinds = [b.kind for b in load_document(str(p)).blocks]
    assert kinds == ["paragraph", "table", "paragraph"], kinds


def test_docx_headings_are_tagged(tmp_path):
    """标题要单独成类 —— 切块级路由靠它做章节递归切。"""
    import docx

    doc = docx.Document()
    doc.add_heading("第一章 总则", level=1)
    doc.add_paragraph("正文。")
    p = tmp_path / "h.docx"
    doc.save(str(p))

    kinds = [b.kind for b in load_document(str(p)).blocks]
    assert kinds == ["heading", "paragraph"], kinds


def test_xlsx_sheets_become_tables(tmp_path):
    p = _make_xlsx(tmp_path / "b.xlsx", sheets={
        "差旅": REIMBURSE_TABLE,
        "考勤": [["项目", "标准"], ["迟到", "30"]],
    })
    doc = load_document(str(p))
    assert doc.parser == "openpyxl"
    caps = [b.meta.get("sheet") for b in doc.tables()]
    assert caps == ["差旅", "考勤"]
    for tb in doc.tables():
        assert missing_numbers(tb.raw_table, tb.text) == []


def test_xlsx_big_sheet_is_split_with_header_repeated(tmp_path):
    """大 sheet 要按行分组，**且每组都要带表头** —— 否则后续组的列名就丢了。"""
    rows = [["行号", "值"]] + [[str(i), str(i * 10)] for i in range(1, 131)]
    p = _make_xlsx(tmp_path / "big.xlsx", sheets={"大表": rows})
    tables = load_document(str(p)).tables()
    assert len(tables) == 3, f"130 行 / 每组 60 行 → 3 组，实际 {len(tables)}"
    for tb in tables:
        assert "行号=" in tb.text, "每组都要重复表头，否则后续组的列名丢失"


def test_csv_becomes_one_table(tmp_path):
    p = tmp_path / "c.csv"
    p.write_text("项目,标准\n迟到,30\n早退,30\n", encoding="utf-8")
    doc = load_document(str(p))
    assert doc.tables() and "标准=30" in doc.tables()[0].text


def test_page_numbers_only_where_meaningful(tmp_path):
    """⚠️ Word/Excel **没有稳定页码** → 必须是 None，不能编一个 0/1。"""
    p = _make_docx(tmp_path / "p.docx", paragraphs=["一段"])
    assert all(b.page is None for b in load_document(str(p)).blocks)


# ---------------------------------------------------------------- 复杂表判定（口径核对点 2）


def test_word_merged_table_is_flagged_complex(tmp_path):
    """⚠️ **口径核对点 2 的回归**：Word 的合并单元格必须被判为复杂表。

    为什么容易漏：python-docx 把合并单元格的文本**重复**到每个被并的格里
    （`差旅标准 | 差旅标准 | 备注`）→ 「行长不齐 / 首行有空」两条判据**都看不出**。
    证据在底层 XML 的 `gridSpan` / `vMerge` 里（口径原话：「查 rowspan/colspan」）。

    漏判的代价不对称：**复杂误判成简单 → 静默污染知识库**（源表该进向量却没进）。
    """
    import docx

    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "差旅标准"
    t.cell(0, 2).text = "备注"
    t.cell(0, 0).merge(t.cell(0, 1))          # 横向合并
    t.cell(1, 0).text = "一线"
    t.cell(1, 1).text = "600"
    t.cell(1, 2).text = "-"
    p = tmp_path / "merged.docx"
    d.save(str(p))

    tables = load_document(str(p)).tables()
    assert len(tables) == 1
    assert tables[0].table_complex is True, "Word 的合并单元格必须判为复杂表"


def test_word_vertical_merge_is_flagged(tmp_path):
    """纵向合并（`vMerge`）同样要判出来 —— 它是另一条 XML 证据。"""
    import docx

    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "部门"
    t.cell(1, 0).text = "IT"
    t.cell(2, 0).text = "IT"
    t.cell(1, 0).merge(t.cell(2, 0))          # 纵向合并
    t.cell(0, 1).text = "人数"
    t.cell(1, 1).text = "3"
    t.cell(2, 1).text = "5"
    p = tmp_path / "vmerge.docx"
    d.save(str(p))

    assert load_document(str(p)).tables()[0].table_complex is True


def test_word_plain_table_is_not_flagged(tmp_path):
    """反向：**普通表不该被误判成复杂**（否则所有表的原表都进向量，等于没判）。"""
    p = _make_docx(tmp_path / "plain.docx", paragraphs=["x"], table=REIMBURSE_TABLE)
    assert load_document(str(p)).tables()[0].table_complex is False


# ---------------------------------------------------------------- 口径核对点 3②


def test_header_and_value_never_cross_rows(tmp_path):
    """口径核对点 3②：**行列头与值不串行**（`Q3` 必须和 `5000` 同句）。

    这条是给"LLM 转"设的闸；**规则转天然满足** —— 但这里仍加一条测试把它钉住，
    因为将来若有人改写渲染逻辑（比如改成"先列所有表头、再列所有值"），
    就会**静默**产生串行，而那正是口径要防的。
    """
    rows = [["季度", "金额"], ["Q3", "5000"], ["Q4", "6000"]]
    nl = render_nl(rows, caption="预算")
    # 每个 值 都必须紧跟在**它自己那行的**列名后面
    assert "季度=Q3；金额=5000" in nl
    assert "季度=Q4；金额=6000" in nl
    # 且不能出现跨行配对
    assert "Q3；金额=6000" not in nl and "Q4；金额=5000" not in nl


def test_parser_layer_has_no_llm_dependency():
    """口径 v1「**不上 LLM**」：解析层（表格转换 + 各格式解析）**不得依赖 LLM 客户端**。

    ⚠️ 用**静态检查**，不是"打桩后再跑一遍" —— 后者是**假保险**：
    若 `tables.py` 写了 `from src.config.litellm_client import complete`，
    再 `monkeypatch.setattr(lc, "complete", boom)` **根本拦不住**
    （那个名字在导入时就已经绑进 `tables` 模块了），测试照样绿。
    """
    import inspect

    from src.document_parser import parsers, tables

    for mod in (tables, parsers):
        src = inspect.getsource(mod)
        for forbidden in ("litellm_client", "get_agent", "acompletion"):
            assert forbidden not in src, (
                f"{mod.__name__} 里出现了 {forbidden!r} —— 解析层不该依赖 LLM"
                f"（口径 v1：规则转换、LLM 调用 0 次、成本 0）"
            )


# ---------------------------------------------------------------- 纯文本 / 编码


def test_plain_text_splits_paragraphs_and_headings(tmp_path):
    p = tmp_path / "n.md"
    p.write_text("# 标题\n\n第一段。\n\n第二段。\n", encoding="utf-8")
    kinds = [b.kind for b in load_document(str(p)).blocks]
    assert kinds == ["heading", "paragraph", "paragraph"], kinds


def test_gbk_file_is_transcoded_not_silently_mangled(tmp_path):
    """非 UTF-8 要**转码**并报警，**不能 errors='ignore' 静默吃字**。"""
    p = tmp_path / "gbk.txt"
    p.write_bytes("这是一段中文，包含数值 12345。".encode("gbk"))
    doc = load_document(str(p))
    assert "12345" in doc.text
    assert any("编码" in w for w in doc.warnings), doc.warnings


# ---------------------------------------------------------------- 警告分层（第1.5批-③）


def test_static_facts_do_not_pollute_warnings(tmp_path):
    """⚠️ **回归**：**静态事实不该进 per-document `warnings`**（第 1 批发现 ③）。

    实测当时 **21/21** 个文件都带 ≥1 条警告，而它们**全是静态的**
    （"该类型只有一个解析器" 21×、"未解析页眉页脚" 6×）→
    **"警告数"这一列再也分不出谁真有毛病**。
    """
    p = _make_docx(tmp_path / "w.docx", paragraphs=["一段"], table=REIMBURSE_TABLE)
    doc = load_document(str(p))
    assert doc.warnings == [], f"docx 不该带静态警告，实际 {doc.warnings}"


def test_route_note_is_still_available_where_it_belongs(tmp_path):
    """去噪**不等于**丢信息：那条路由事实仍在 `RoutePlan.note` 上，需要的人能看到。"""
    p = _make_docx(tmp_path / "w.docx", paragraphs=["一段"])
    assert route_file(str(p)).note, "路由的 note 不该被一起删掉"


def test_dynamic_warnings_still_fire(tmp_path):
    """**反向**：别为了去噪把真闸也拆了 —— 真·动态问题必须继续报。"""
    p = tmp_path / "gbk.txt"
    p.write_bytes("中文，含数值 12345。".encode("gbk"))
    assert any("编码" in w for w in load_document(str(p)).warnings)


# ⚠️ 这里**故意没有再写**一条"PDF 空页仍报警"的测试：
# 造一个"有页但无文本层"的 PDF 成本高，而写一条"构造 ParsedDoc 再断言它的字段"
# 是**恒真的废话** —— 它看起来像覆盖、其实什么都没测（同"假绿灯"那类问题）。
# `test_dynamic_warnings_still_fire` 已覆盖"动态警告机制仍活着"这件事。


def test_markdown_code_fence_is_a_code_block(tmp_path):
    p = tmp_path / "code.md"
    p.write_text("正文。\n\n```python\nprint(1)\n```\n", encoding="utf-8")
    kinds = [b.kind for b in load_document(str(p)).blocks]
    assert "code" in kinds


# ---------------------------------------------------------------- 代码文件（2.0-24 前半）


def test_code_file_becomes_one_code_block(tmp_path):
    """**回归**：代码文件必须产出**整块** `code` 块。

    以前 `.py` 走"按空行分段" → 实测 `23_edge_collector.py` 被切成 **20 块**，
    而 `chunking` 的「代码整块」策略只对 `kind == "code"` 生效 →
    **那条策略对代码文件从未生效**（声明了、却没作用）。
    """
    p = tmp_path / "m.py"
    p.write_text("import os\n\n\ndef f():\n    return 1\n\n\nclass C:\n    pass\n", encoding="utf-8")
    doc = load_document(str(p))
    assert doc.kind_counts() == {"code": 1}
    assert doc.blocks[0].meta.get("lang") == "py"


def test_code_block_is_never_split(tmp_path):
    """整块不切：`chunk_size` 压到极小也**必须**只有一块。"""
    from src.document_parser.chunking import split_blocks

    p = tmp_path / "big.py"
    p.write_text("\n\n\n".join(f"def f{i}():\n    return {i}" for i in range(50)), encoding="utf-8")
    cs = split_blocks(load_document(str(p)), chunk_size=10, overlap=2)
    assert len(cs) == 1, "代码块整块不切（tree-sitter 按函数切是 2.0-24 后半，本轮不做）"


def test_code_extensions_have_one_source_of_truth():
    """⚠️ `.py` 是不是代码，必须来自**路由表** —— 不能另抄一份清单（两份必然漂移）。"""
    from src.document_parser.routing import kind_by_extension, supported_extensions

    assert kind_by_extension("py") == "code"
    assert kind_by_extension("PY") == "code"          # 大小写不敏感
    assert kind_by_extension(".py") == "code"         # 带点也认
    assert kind_by_extension("xlsx") == "xlsx"
    assert "py" in supported_extensions()


def test_same_content_in_txt_is_not_a_code_block(tmp_path):
    """反向：**同样内容**放在 `.txt` 里**不该**变成 code 块（证明派发真的按扩展名）。"""
    body = "import os\n\n\ndef f():\n    return 1\n"
    p = tmp_path / "m.txt"
    p.write_text(body, encoding="utf-8")
    # 断言抓**语义**（"不该有 code 块"），不数它切成几段 ——
    # 段数取决于按空行切的实现细节，写死会把无关改动弄红（今天第二次踩这个）
    assert "code" not in load_document(str(p)).kind_counts()
