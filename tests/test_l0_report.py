"""L0 体检脚本的冒烟测试。

为什么值得测：`2.0-39` 把它定成了**纪律**（改 `chunk_size` 前必须附 L0 报告），
所以它**是一道闸** —— 而闸自己坏掉是不会有人发现的（本项目反复踩的那类）。

用**子进程**跑真实入口（而不是 import 内部函数）：
`eval/` 不是包，且这样测的正是"用户实际怎么用它"。
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "eval" / "l0_report.py"


def _make_corpus(tmp: Path) -> Path:
    """造一个最小语料：1 个带表格的 docx + 1 个 txt。"""
    import docx

    corpus = tmp / "corpus"
    corpus.mkdir()
    d = docx.Document()
    d.add_paragraph("本制度适用于全体员工。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "城市"
    t.cell(0, 1).text = "住宿上限"
    t.cell(1, 0).text = "一线"
    t.cell(1, 1).text = "600"
    d.save(str(corpus / "a.docx"))
    (corpus / "b.txt").write_text("第一段。\n\n第二段。\n", encoding="utf-8")
    return corpus


def test_l0_report_runs_and_reports(tmp_path):
    corpus = _make_corpus(tmp_path)
    out = tmp_path / "l0.json"
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(corpus), "--json", str(out)],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0, f"脚本失败：{r.stderr[-400:]}"

    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["summary"]["files"] == 2
    assert data["summary"]["ok"] == 2
    # ⭐ 这两条是"报告可信"的前提：表被认出来了，且数值保真没失败
    assert data["summary"]["tables"] == 1, "docx 里的表格必须被计到"
    assert data["summary"]["table_missing_numbers"] == 0
    # 逐文件明细要带上"**实际生效的解析器**"（协议 P-2）
    by_name = {f["file"]: f for f in data["files_detail"]}
    assert by_name["a.docx"]["parser"] == "python-docx"
    assert by_name["b.txt"]["parser"] == "plain"


def test_l0_report_survives_a_broken_file(tmp_path):
    """⚠️ **一个坏文件不该中断整批** —— 它要进"失败清单"，其余照常体检。

    （对照 `eval/ingest.py` 曾经"一篇坏就中断整批"那个坑。）
    """
    corpus = _make_corpus(tmp_path)
    (corpus / "c.html").write_text("<html><body>x</body></html>", encoding="utf-8")
    out = tmp_path / "l0.json"
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(corpus), "--json", str(out)],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0, "有坏文件时不该非零退出"
    s = json.loads(out.read_text(encoding="utf-8"))["summary"]
    assert s["ok"] == 2 and s["failed"] == 1, f"应 2 成功 / 1 失败，实际 {s}"
    detail = json.loads(out.read_text(encoding="utf-8"))["files_detail"]
    # 失败**原因**要可读（不是一句"失败了"）
    assert any("HTML" in (f.get("error") or "") for f in detail), "失败原因要可读"


def test_l0_report_rejects_a_missing_dir(tmp_path):
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(tmp_path / "nope")],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 1, "目录不存在要非零退出（别静默报 0 个文件）"


@pytest.mark.parametrize("flag", ["--no-detail"])
def test_no_detail_flag_works(tmp_path, flag):
    corpus = _make_corpus(tmp_path)
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(corpus), flag,
         "--json", str(tmp_path / "l0.json")],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0
    assert "【汇总】" in r.stdout, "汇总必须仍然打印"


# ---------------------------------------------------------------- 2.0-51


def test_pct_never_leaves_the_sample_range():
    """`2.0-51①`：分位数**必须落在样本区间内** —— 回归闸。

    旧实现用 `statistics.quantiles`（默认 `exclusive`），**小样本会外推**：
    实测给过 `p99 = 2932 > max = 2259`，以及把 `[192, 2259]` 的 p50 报成 **1225**
    （那是**两点的中点**，不是分位数）。这两个都是真踩过的数，所以直接拿来当用例。
    """
    from eval.l0_report import pct_nearest_rank

    for vals in ([192, 2259], [637, 851], [120, 213, 256, 933], [86, 219, 1406]):
        for q in (50, 90, 99):
            v = pct_nearest_rank(vals, q)
            assert min(vals) <= v <= max(vals), f"{vals} 的 p{q} = {v} 超出样本区间"
    # 反向用例：旧口径在这里给 1225（中点）
    assert pct_nearest_rank([192, 2259], 50) == 192, "nearest-rank 的 p50 应取到实际样本值 192"


def test_caption_count_is_chunk_level(tmp_path):
    """`2.0-51②`：caption 计数必须发生在**切块后** —— 否则测不到 `2.0-46`/`2.0-48` 的机制。

    ⚠️ 夹具刻意让**加粗手工标题夹在两张表之间**（标题因此独自成一个 text 段），
    这样 `chunking._merge_heading_captions` 才会把它并成表格的 caption。
    块级（`doc.tables()`）的 text **没有** `表「` 前缀 —— 所以旧口径会数成 0。
    """
    import docx

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    d = docx.Document()
    d.add_paragraph("本制度适用于全体员工。")
    for label in ("前置项", "营业成本"):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "项目"
        t.cell(0, 1).text = "金额"
        t.cell(1, 0).text = label
        t.cell(1, 1).text = "16930"
        if label == "前置项":
            d.add_paragraph().add_run("三、成本与费用").bold = True
    d.save(str(corpus / "a.docx"))

    out = tmp_path / "l0.json"
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(corpus), "--json", str(out)],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0, f"脚本失败：{r.stderr[-400:]}"
    s = json.loads(out.read_text(encoding="utf-8"))["summary"]
    assert s["tables"] == 2
    assert s["tables_with_caption"] == 1, (
        f"加粗标题应给紧邻的表补上 caption（切块层）；实际 {s['tables_with_caption']}"
    )
    assert s["tables_without_caption"] == 1
    # 原始块长要能复核（2.0-51① 的配套要求）
    detail = json.loads(out.read_text(encoding="utf-8"))["files_detail"]
    assert all("lens_asc" in f for f in detail if not f.get("error"))


def test_pct_has_a_single_source():
    """`2.0-51③`：分位实现**只能有一份** —— 评测脚本必须 import，不许各写一份。

    两份实现现在一致，但**两份就是"同一事实两个来源"**，而漂移是**静默**的
    （本项目已因同类事故吃过两次：`upload_api._ALLOWED` 的两份清单 /
    `parse_plain` 的两处扩展名判断）。

    ⚠️ 断言用 `is`（**同一个函数对象**），不是"结果相等" ——
    有人**复制一份算法一样**的实现，这条也必须红。
    """
    import eval.run_excel_numeric_eval as ev
    from eval.l0_report import pct_nearest_rank

    assert ev.pct_nearest_rank is pct_nearest_rank, (
        "评测脚本必须 import L0 那份分位实现，不许自带一份"
    )
