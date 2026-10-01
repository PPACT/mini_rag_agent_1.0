"""`2.0-48` 探针的回归闸。

为什么值得测：这是**"先量"的依据** —— 它给出的数字会被拿去定判据。
**量错了比不量更糟**：会拿一个错误的基数去裁决"要不要在解析层上启发式"。

⭐ **反向用例是必须的**（`dev-agent.md` 纪律）：只测"能捞出来"，
**测不出"把正文当成标题捞出来"** —— 而误报正是本项最贵的失败模式
（同"误判是把章节标题错当成 caption"那类污染）。

⚠️ 最该盯住的一条：**"正文基准字号"必须真的来自正文**。
实测踩过：不读 `docDefaults` 时，正文全解析成 `None`，
基准只能从**少数显式带字号的段落（恰恰就是标题）**里算出
→ **基准 = 拿标题当正文** → 整档数字失真。`test_baseline_comes_from_body` 就是钉它。
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "eval" / "probe_docx_heading.py"
sys.path.insert(0, str(ROOT))

from eval.probe_docx_heading import _scan_docx  # noqa: E402


def _docx_with_manual_title(tmp: Path, name: str = "a.docx") -> Path:
    """一个"手工格式化标题"的 docx：正文用默认字号，标题靠**加粗 + 大字号**。

    这正是 `2.0-48` 要量的对象 —— 段落是 `Normal` 样式，**没被 Word 标成 Heading**。
    """
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.add_paragraph("本规定自发布之日起施行，适用于全体研发人员。")      # 正文：默认字号
    p = d.add_paragraph()
    r = p.add_run("三、成本与费用")                                        # 手工标题
    r.bold = True
    r.font.size = Pt(16)
    d.add_paragraph("研发费用按项目归集，月末统一结转。")                  # 正文
    path = tmp / name
    d.save(str(path))
    return path


def test_baseline_comes_from_body(tmp_path):
    """⭐ 核心回归闸：**基准字号来自正文（默认 11pt），不是标题的 16pt**。

    若不读 `docDefaults`，正文的 `size_pt` 全为 `None` → 基准被标题的 16pt 顶上去。
    """
    rec = _scan_docx(_docx_with_manual_title(tmp_path), chunk_size=512, overlap=64)
    assert rec["baseline_size_pt"] == 11.0, (
        f"基准应为正文默认 11pt，实际 {rec['baseline_size_pt']} —— "
        "疑似 docDefaults 没接进继承链（基准被标题污染）"
    )
    assert rec["n_size_unresolved"] == 0, "有段落字号解析不出来 → 基准不可信"


def test_flags_manual_title(tmp_path):
    """正文准的标题要被捞出来，且**标注它是非 heading 的手工标题**。

    ⚠️ 这里用 `chunk_size=20` **刻意把每段拆成独立块** ——
    512 时三段会被并成一块（48 字符），标题就**根本不进极短块清单**。
    （这也顺带印证：判据的对象是**块**不是段落，所以候选段 ≠ 极短块。）
    """
    rec = _scan_docx(_docx_with_manual_title(tmp_path), chunk_size=20, overlap=0)
    cands = [p["text"] for p in rec["cand_bold_short"]]
    assert "三、成本与费用" in cands, f"手工标题没被捞出来：{cands}"
    assert rec["n_heading"] == 0, "本文件不应有样式级 heading（这正是缺口所在）"
    # `2.0-46` 残留极短块的交叉核对：该标题应**自己成块**且被判为"可捞回"
    tiny = [t for t in rec["tiny_chunks"] if t["text"] == "三、成本与费用"]
    assert tiny and tiny[0]["bold_all"] is True, "极短块交叉核对没对上"
    # 反向面：**多段合并**的短块拿不到单一格式事实 → 必须是 `None`，**不能**瞎猜成 True
    merged = [t for t in rec["tiny_chunks"] if not t["single_para"]]
    assert all(t["bold_all"] is None for t in merged), "合并不该伪造格式事实"


def test_plain_body_is_not_flagged(tmp_path):
    """⚠️ **反向用例**：短但不加粗、字号正常的正文，**不许**被当成标题捞出来。"""
    import docx

    d = docx.Document()
    d.add_paragraph("本条适用于全体人员。")            # 短正文（10 字）
    d.add_paragraph("研发费用按项目归集，月末统一结转。")
    path = tmp_path / "b.docx"
    d.save(str(path))

    rec = _scan_docx(path, chunk_size=512, overlap=64)
    assert rec["c1_bold"] == 0, "没有加粗段，不该有任何 C1 候选"
    assert rec["d2_big_short"] == 0, "没有大字号段，不该有任何 D2 候选"


def test_styled_heading_is_excluded(tmp_path):
    """已认出的 heading **不许**再进候选 —— 否则同一事实有了两个来源。

    反向用例的另一面：如果候选里混进了 `Heading 1`，说明"分母"取错了。
    """
    import docx

    d = docx.Document()
    d.add_paragraph("正文一句话，用来撑出正文基准字号。")
    h = d.add_paragraph("一、总则", style="Heading 1")
    h.runs[0].bold = True
    path = tmp_path / "c.docx"
    d.save(str(path))

    rec = _scan_docx(path, chunk_size=512, overlap=64)
    assert rec["n_heading"] == 1, "样式 heading 要被认出来"
    assert "一、总则" not in [p["text"] for p in rec["cand_bold_short"]], (
        "样式 heading 不该出现在候选里（它是已认定的事实，不是缺口）"
    )


def test_cli_runs_and_writes_json(tmp_path):
    """CLI 冒烟：跑真实入口，产物 JSON 结构要齐（与 `l0_report` 同风格）。"""
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _docx_with_manual_title(corpus)
    out = tmp_path / "probe.json"
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(corpus), "--json", str(out)],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0, f"脚本失败：{r.stderr[-400:]}"
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["summary"]["files"] == 1
    assert "全量明细" in r.stdout
    detail = {f["file"]: f for f in data["files_detail"]}
    assert detail["a.docx"]["doc_defaults"]["size_pt"] == 11.0


def test_cli_rejects_missing_dir(tmp_path):
    r = subprocess.run(
        [sys.executable, str(SCRIPT), "--dir", str(tmp_path / "nope")],
        capture_output=True, text=True, encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 1, "目录不存在要非零退出（别静默报 0 个文件）"
