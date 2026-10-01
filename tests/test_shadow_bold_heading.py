"""`2.0-48` 判据（整段加粗）与 shadow 台子的回归闸。

⚠️ **反向用例是重点**（文档侧 `交流区 §1.40④③`：**"找不到反例也要有机制"**）——
"整段加粗但不是标题"的**强调句**会被这条判据误认，这是它的**固有代价**。
所以本文件**不假装它不存在**，而是把它**钉成可复现的事实**：
一旦将来有人加了一条"长度闸"把误认挡住了，`test_known_failure_mode...` 会红 ——
**那时候要连带更新口径**，而不是让它悄悄改变行为。

⚠️ 另一条同样重要：**开关必须仍然可切回去**。`2.0-48` 开启后
`parse_docx(path)` 的默认已是新判据，但**显式传 `_is_heading` 必须仍能得到旧行为** ——
否则 `eval/shadow_bold_heading.py` 的 A/B 台子第二天就失效了（两边同判据、差值恒为 0）。
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from eval.shadow_bold_heading import _looks_like_title  # noqa: E402
from src.document_parser.parsers import (  # noqa: E402
    _is_heading,
    heading_rule_style_or_bold,
    parse_docx,
)
from src.document_parser.routing import parse_document  # noqa: E402
from src.document_parser.table_recovery import recover_tables  # noqa: E402


def _add_table(d, label: str, value: str) -> None:
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "金额"
    t.cell(1, 0).text = label
    t.cell(1, 1).text = value


def _docx(tmp: Path, name: str = "a.docx", *, bold_title: str | None = None,
          styled_heading: str | None = None, table: bool = True,
          bold_sentence: str | None = None) -> Path:
    """⚠️ `table=True` 时**标题两侧都放表** —— 因为 `_segments` 会把**连续的非原子块并成一段**，
    只有「前后都是原子块」时标题才**独自成段**，才可能触发 `2.0-46` 的 caption 合并。
    这正是真实语料里 `三、成本与费用` 的形态；夹具不对，测的就是另一回事。
    """
    import docx

    d = docx.Document()
    d.add_paragraph("本制度适用于全体员工，自发布之日起施行。")
    if bold_sentence:
        d.add_paragraph().add_run(bold_sentence).bold = True
    if table:
        _add_table(d, "前置项", "1")
    if styled_heading:
        d.add_paragraph(styled_heading, style="Heading 1")
    if bold_title:
        d.add_paragraph().add_run(bold_title).bold = True
    if table:
        _add_table(d, "营业成本", "16930")
    path = tmp / name
    d.save(str(path))
    return path


_UNSET = object()


def _kind_of(path: Path, rule=_UNSET) -> dict[str, str]:
    """`{段落文本: kind}` —— 便于按文本断言。

    ⚠️ **不传 `rule` = 走生产默认**（不显式传参）—— 这样"默认变没变"本身也是被测的。
    """
    kw = {} if rule is _UNSET else {"heading_rule": rule}
    return {b.text: b.kind for b in parse_docx(str(path), **kw).blocks}


# ---------------------------------------------------------------- 判据本体


def test_default_rule_is_the_new_one(tmp_path):
    """⭐ `2.0-48` 已开启：`parse_docx` **默认**就认加粗（`交流区 §1.42②`）。"""
    path = _docx(tmp_path, bold_title="三、成本与费用")
    kinds = _kind_of(path)
    assert kinds.get("三、成本与费用") == "heading", "默认应已是新判据（样式 OR 加粗）"


def test_old_rule_still_available(tmp_path):
    """⭐ **开关仍能切回旧判据** —— 显式传 `_is_heading` 时，加粗标题**仍**是 paragraph。

    这条是 `eval/shadow_bold_heading.py` **A/B 台子的地基**：
    一红就说明台子两边变成同一个判据了（差值恒为 0，验不出任何东西）。
    """
    path = _docx(tmp_path, bold_title="三、成本与费用")
    kinds = _kind_of(path, _is_heading)
    assert kinds.get("三、成本与费用") == "paragraph", (
        "显式传旧判据时必须退回只认样式名 —— 否则 A/B 台子失效"
    )


def test_production_entry_uses_new_rule(tmp_path):
    """⭐ **验收本体**：生产入口 `parse_document`（不是 `parse_docx`）也得认加粗。

    只测 `parse_docx` 不够 —— 接线断在 `routing` 那一层的话，上面几条**照样绿**。
    """
    path = _docx(tmp_path, bold_title="三、成本与费用")
    kinds = {b.text: b.kind for b in parse_document(str(path)).blocks}
    assert kinds.get("三、成本与费用") == "heading", "生产入口没接上新判据"


def test_bold_headings_are_recognized_under_new_rule(tmp_path):
    path = _docx(tmp_path, bold_title="三、成本与费用")
    kinds = _kind_of(path, heading_rule_style_or_bold)
    assert kinds.get("三、成本与费用") == "heading"


def test_styled_heading_not_lost(tmp_path):
    """⚠️ 新判据**不能顶掉**已认定的事实（样式标题）—— 否则会倒退。"""
    path = _docx(tmp_path, styled_heading="一、总则")
    kinds = _kind_of(path, heading_rule_style_or_bold)
    assert kinds.get("一、总则") == "heading"


def test_plain_short_paragraph_is_not_a_heading(tmp_path):
    """反向面：**不加粗**的短段落不该被认成标题（判据不是"什么都认"）。"""
    path = _docx(tmp_path)          # 没有加粗标题
    kinds = _kind_of(path, heading_rule_style_or_bold)
    assert kinds.get("本制度适用于全体员工，自发布之日起施行。") == "paragraph"


def test_known_failure_mode_bold_emphasis_sentence(tmp_path):
    """⭐ **已知失败模式（反向用例）**：整段加粗的**强调句**会被误认成标题。

    **这是本判据的固有代价，不是 bug** —— `2.0-48` 先 shadow 就是为了拿数字决定
    要不要接受它（文档侧 `§1.40④③`：找不到反例也要有机制）。

    ⚠️ 若将来加了"长度闸"或"句读闸"把这条挡住了，**本测试会红** ——
    那是**行为变了**，要连着口径一起更新，别让它静默漂移。
    """
    sent = ("注意：本节涉及的全部金额均为含税口径，与第四节的口径不同，"
            "跨节比较前请先按增值税率折算，切勿直接混用两份数据。")
    path = _docx(tmp_path, bold_sentence=sent)

    kinds = _kind_of(path, heading_rule_style_or_bold)
    assert kinds.get(sent) == "heading", "判据只看加粗 —— 强调句会被认成标题（预期行为）"

    # 而"像标题"这个**展示用代理指标**应当把它挡在门外 —— 两者差值就是误认数
    assert _looks_like_title(sent) is False, (
        "长且以句读结尾 → 代理指标不该算它像标题；"
        "若这里变了，shadow 报告的'像标题'列口径也变了"
    )
    assert len(sent) > 40, "本用例要真是'长句'，否则证明不了长度能区分"


# ---------------------------------------------------------------- 切块后果


def test_bold_title_before_table_gains_caption(tmp_path):
    """⭐ 判据的**唯一 content 改动路径**：紧邻标题的表格补上 caption。

    这也是 `2.0-46` 的机制被"接上"的地方 —— 在此之前 docx 分支够不着它。
    """
    from src.config.settings import get_settings
    from src.document_parser.chunking import split_blocks

    s = get_settings()
    path = _docx(tmp_path, bold_title="三、成本与费用")
    shadow = split_blocks(
        parse_docx(str(path), heading_rule=heading_rule_style_or_bold),
        s.chunk_size, s.chunk_overlap)
    captioned = [c for c in shadow if c.text.startswith("表「")]
    assert captioned, f"表格应拿到 caption，实际块：{[c.text[:24] for c in shadow]}"
    assert captioned[0].text.startswith("表「三、成本与费用」：")
    assert captioned[0].title == "三、成本与费用"


def test_old_rule_has_no_caption(tmp_path):
    """对照组：**旧判据**下同一文件没有 caption —— 证明上一条的改变来自判据本身。"""
    from src.config.settings import get_settings
    from src.document_parser.chunking import split_blocks

    s = get_settings()
    path = _docx(tmp_path, bold_title="三、成本与费用")
    base = split_blocks(parse_docx(str(path), heading_rule=_is_heading),
                        s.chunk_size, s.chunk_overlap)
    assert not any(c.text.startswith("表「") for c in base)


# ---------------------------------------------------------------- AB 台子


def test_harness_matches_production_pipeline(tmp_path):
    """⭐ **AB 台子自检**：`parse_docx` + 同一后处理 必须等于生产入口 `parse_document`。

    实测踩过：漏掉 `recover_tables` → 伪表格被当普通段落 → **AB 两组差了两个变量**（P-9），
    整批数字失真。这条把那个坑钉住。
    """
    path = _docx(tmp_path, bold_title="三、成本与费用")
    replay, _ = recover_tables(parse_docx(str(path)))
    prod = parse_document(str(path))
    assert [(b.kind, b.text) for b in replay.blocks] == \
           [(b.kind, b.text) for b in prod.blocks], "复算链与生产链不一致"
